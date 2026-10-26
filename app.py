import os
from datetime import date, timedelta

import pandas as pd
import streamlit as st
from openpyxl import load_workbook

st.set_page_config(page_title="Internship Task Management", page_icon="📋", layout="wide")
st.title("📋 Internship Task Management Dashboard")

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
FILE = os.path.join(BASE_DIR, "Internship_Task_Management.xlsx")
SHEET = "Intern Details"

if not os.path.exists(FILE):
    st.error("Excel file not found. Keep app.py and Internship_Task_Management.xlsx in the same folder.")
    st.stop()

try:
    df = pd.read_excel(FILE, sheet_name=SHEET)
except Exception as exc:
    st.error(f"Unable to read Excel file: {exc}")
    st.stop()

df.columns = df.columns.astype(str).str.strip()
columns = [
    "Intern ID", "Intern Name", "Email ID", "Task ID", "Task Name",
    "Assigned Date", "Review Date", "Review Number", "Status", "Reviewer", "Feedback"
]
for col in columns:
    if col not in df.columns:
        df[col] = ""

email_series = df["Email ID"].fillna("").astype(str).str.strip()
total_interns = email_series[email_series.ne("")].str.lower().nunique()
task_names = df["Task Name"].fillna("").astype(str).str.strip()
statuses = df["Status"].fillna("").astype(str).str.strip().str.lower()
assigned = int(task_names.ne("").sum())
pending = int(statuses.isin(["pending", "under review"]).sum())
completed = int(statuses.eq("completed").sum())

c1, c2, c3, c4 = st.columns(4)
c1.metric("Total Interns", total_interns)
c2.metric("Tasks Assigned", assigned)
c3.metric("Pending Reviews", pending)
c4.metric("Completed Tasks", completed)
st.divider()


def save_cells(row_index, fields):
    wb = load_workbook(FILE)
    ws = wb[SHEET]
    headers = {str(cell.value).strip(): cell.column for cell in ws[1] if cell.value is not None}
    excel_row = int(row_index) + 2
    for field in fields:
        if field in headers:
            value = df.at[row_index, field]
            if pd.isna(value):
                value = None
            ws.cell(row=excel_row, column=headers[field]).value = value
    wb.save(FILE)
    wb.close()

st.subheader("➕ Assign New Task")
valid_emails = sorted(set(email_series[email_series.ne("")].tolist()))
if valid_emails:
    with st.form("assign_task"):
        email = st.selectbox("Intern Email", valid_emails)
        task_id = st.text_input("Task ID")
        task_name = st.text_input("Task Name")
        assigned_date = st.date_input("Assigned Date", value=date.today())
        reviewer = st.text_input("Reviewer Name")
        submitted = st.form_submit_button("Assign Task")

    if submitted:
        if not task_id.strip() or not task_name.strip():
            st.error("Please enter both Task ID and Task Name.")
        else:
            matches = df.index[df["Email ID"].fillna("").astype(str).str.strip() == email].tolist()
            if not matches:
                st.error("Intern email not found.")
            else:
                row = matches[0]
                df.at[row, "Task ID"] = task_id.strip()
                df.at[row, "Task Name"] = task_name.strip()
                df.at[row, "Assigned Date"] = assigned_date
                df.at[row, "Review Date"] = assigned_date + timedelta(days=1)
                df.at[row, "Review Number"] = 1
                df.at[row, "Status"] = "Pending"
                df.at[row, "Reviewer"] = reviewer.strip()
                df.at[row, "Feedback"] = ""
                try:
                    save_cells(row, ["Task ID", "Task Name", "Assigned Date", "Review Date", "Review Number", "Status", "Reviewer", "Feedback"])
                    st.success("Task assigned successfully!")
                    st.rerun()
                except Exception as exc:
                    st.error(f"Could not save task. Close the Excel file and try again. Details: {exc}")
else:
    st.warning("No intern email addresses found in the Excel file.")

st.divider()
st.subheader("📑 Intern Task Details")
st.dataframe(df[columns], use_container_width=True, hide_index=True)

st.subheader("🔍 Search Intern")
search_email = st.text_input("Enter intern email to search")
if search_email.strip():
    result = df[df["Email ID"].fillna("").astype(str).str.strip().str.lower() == search_email.strip().lower()]
    if result.empty:
        st.warning("Intern not found.")
    else:
        st.dataframe(result[columns], use_container_width=True, hide_index=True)

st.divider()
st.subheader("📝 Update Review")
assigned_rows = df.index[task_names.ne("")].tolist()
if assigned_rows:
    row_index = st.selectbox(
        "Select task",
        assigned_rows,
        format_func=lambda i: f"{df.at[i, 'Intern Name']} | {df.at[i, 'Task Name']} | {df.at[i, 'Email ID']}"
    )
    with st.form("update_review"):
        status = st.selectbox("Review Status", ["Pending", "Under Review", "Completed", "Rejected"])
        feedback = st.text_area("Reviewer Feedback")
        next_review = st.date_input("Next Review Date", value=date.today() + timedelta(days=1))
        update = st.form_submit_button("Update Review")

    if update:
        df.at[row_index, "Status"] = status
        df.at[row_index, "Feedback"] = feedback
        fields = ["Status", "Feedback"]
        if status in ["Rejected", "Under Review"]:
            current = df.at[row_index, "Review Number"]
            try:
                current = int(current)
            except (ValueError, TypeError):
                current = 1
            df.at[row_index, "Review Number"] = current + 1
            df.at[row_index, "Review Date"] = next_review
            fields.extend(["Review Number", "Review Date"])
        try:
            save_cells(row_index, fields)
            st.success("Review updated successfully!")
            st.rerun()
        except Exception as exc:
            st.error(f"Could not update review. Close the Excel file and try again. Details: {exc}")
else:
    st.info("No tasks assigned yet.")
